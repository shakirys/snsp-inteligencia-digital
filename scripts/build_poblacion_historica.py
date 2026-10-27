#!/usr/bin/env python3
"""
scripts/build_poblacion_historica.py
-----------------------------------------------------------------------
Población histórica SUIVE 2017-2026 (ACT26) — NUEVO, aditivo.

Entradas (no se modifican):
  data/raw/BASE_POBLACION_SUIVE_QRO_2017_2026_PLANA.xlsx  (hoja POBLACION_BASE)
  data/geo/Municipios_Queretaro_18.topojson               (límites municipales)
  data/geo/Municipios_Queretaro_18.geojson                (validación cruzada)
  data/raw/GEOGRAFIA_QUERETARO_POWER_BI.xlsx              (catálogo DIM_GEOGRAFIA_MUNICIPAL)
  data/raw/EXTRACCION_POBLACION_SUIVE_QRO_2022_2026.xlsx  (conciliación con Morbilidad)

Salidas:
  data/real/poblacion_historica_2017_2026.js   -> window.SNSP_POBLACION_HISTORICA_DATA
  data/real/geo_municipios_queretaro.js        -> window.SNSP_GEO_QRO_MUNICIPIOS
  data/processed/poblacion_historica_2017_2026.json
  data/processed/auditoria_poblacion_historica.json

Reglas:
  * Los valores de población se conservan con su precisión ORIGINAL
    (incluye decimales). Se agregan sin redondear y sólo se redondean al
    mostrar. No se redondea ninguna fila antes de sumar.
  * La relación base <-> mapa <-> catálogo se hace por CLAVE de municipio
    (001-018 / CVEGEO 22xxx), nunca por texto: los nombres difieren entre
    fuentes. Si alguna clave o jurisdicción no coincide, el script ABORTA.
  * Se respetan los 11 grupos de edad tal como vienen (amplitudes
    desiguales). No se generan quinquenales.
  * La fuente se describe como "SUIVE — base de población 2017-2026". No
    se atribuye a CONAPO hasta confirmar su fuente metodológica.
  * Fuente distinta e independiente del padrón ASIS 2025 por localidad
    (data/real/poblacion_2025.js), que no se toca.
-----------------------------------------------------------------------
"""
import json
import math
import os
import sys
from collections import defaultdict
from datetime import date

import openpyxl

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
RAW = os.path.join(ROOT, "data/raw/BASE_POBLACION_SUIVE_QRO_2017_2026_PLANA.xlsx")
TOPO = os.path.join(ROOT, "data/geo/Municipios_Queretaro_18.topojson")
GEOJSON = os.path.join(ROOT, "data/geo/Municipios_Queretaro_18.geojson")
PBI = os.path.join(ROOT, "data/raw/GEOGRAFIA_QUERETARO_POWER_BI.xlsx")
MORB = os.path.join(ROOT, "data/raw/EXTRACCION_POBLACION_SUIVE_QRO_2022_2026.xlsx")
OUT_JS = os.path.join(ROOT, "data/real/poblacion_historica_2017_2026.js")
OUT_GEO_JS = os.path.join(ROOT, "data/real/geo_municipios_queretaro.js")
OUT_JSON = os.path.join(ROOT, "data/processed/poblacion_historica_2017_2026.json")
OUT_AUDIT = os.path.join(ROOT, "data/processed/auditoria_poblacion_historica.json")

# 11 grupos tal como vienen en la fuente (orden de menor a mayor edad).
GRUPOS = [
    ("De 00 años", "0 años", 1),
    ("De 01 a 04 años", "1–4", 4),
    ("De 05 a 09 años", "5–9", 5),
    ("De 10 a 14 años", "10–14", 5),
    ("De 15 a 19 años", "15–19", 5),
    ("De 20 a 24 años", "20–24", 5),
    ("De 25 a 44 años", "25–44", 20),
    ("De 45 a 49 años", "45–49", 5),
    ("De 50 a 59 años", "50–59", 10),
    ("De 60 a 64 años", "60–64", 5),
    ("De 65 y más años", "65 y más", None),  # abierto
]
SEXOS = {"Masculino": 0, "Femenino": 1}
NOMBRE_JS = {1: "Querétaro", 2: "San Juan del Río", 3: "Cadereyta", 4: "Jalpan"}

errores = []


def fail(msg):
    errores.append(msg)
    print("ERROR:", msg)


# ---------------------------------------------------------------------
# 1. Base de población
# ---------------------------------------------------------------------
def leer_base():
    wb = openpyxl.load_workbook(RAW, read_only=True, data_only=True)
    if "POBLACION_BASE" not in wb.sheetnames:
        sys.exit("No existe la hoja POBLACION_BASE")
    ws = wb["POBLACION_BASE"]
    rows = list(ws.iter_rows(values_only=True))
    header = [str(h).strip() if h is not None else "" for h in rows[0]]
    esperado = ["Entidad", "Año", "Jurisdicción", "Municipio", "Sexo", "Grupo de edad", "Poblacion"]
    if header != esperado:
        sys.exit(f"Encabezado inesperado: {header}")
    return [dict(zip(header, r)) for r in rows[1:] if any(v is not None for v in r)]


# ---------------------------------------------------------------------
# 2. TopoJSON -> SVG con topología compartida (bordes sin huecos)
# ---------------------------------------------------------------------
def decode_arcs(topo):
    sx, sy = topo["transform"]["scale"]
    tx, ty = topo["transform"]["translate"]
    arcs = []
    for arc in topo["arcs"]:
        x = y = 0
        pts = []
        for dx, dy in arc:
            x += dx
            y += dy
            pts.append((x * sx + tx, y * sy + ty))
        arcs.append(pts)
    return arcs


def dp(points, tol):
    """Douglas-Peucker; conserva extremos (los nodos compartidos)."""
    if len(points) < 3:
        return points
    keep = [False] * len(points)
    keep[0] = keep[-1] = True
    stack = [(0, len(points) - 1)]
    while stack:
        a, b = stack.pop()
        ax, ay = points[a]
        bx, by = points[b]
        dx, dy = bx - ax, by - ay
        norm = math.hypot(dx, dy)
        dmax, idx = -1, -1
        for i in range(a + 1, b):
            px, py = points[i]
            d = abs(dy * px - dx * py + bx * ay - by * ax) / norm if norm else math.hypot(px - ax, py - ay)
            if d > dmax:
                dmax, idx = d, i
        if dmax > tol and idx > 0:
            keep[idx] = True
            stack.append((a, idx))
            stack.append((idx, b))
    return [p for p, k in zip(points, keep) if k]


def construir_geo(topo, tol=0.0006, width=620):
    arcs = [dp(a, tol) for a in decode_arcs(topo)]
    minx, miny, maxx, maxy = topo["bbox"]
    lat_mid = math.radians((miny + maxy) / 2)
    kx = math.cos(lat_mid)
    scale = width / ((maxx - minx) * kx)
    height = (maxy - miny) * scale
    pad = 8

    def proj(p):
        return ((p[0] - minx) * kx * scale + pad, (maxy - p[1]) * scale + pad)

    def ring(idx_list):
        pts = []
        for i in idx_list:
            a = arcs[i] if i >= 0 else list(reversed(arcs[~i]))
            pts.extend(a if not pts else a[1:])
        return pts

    feats = []
    for g in topo["objects"]["municipios"]["geometries"]:
        polys = g["arcs"] if g["type"] == "MultiPolygon" else [g["arcs"]]
        d_parts, area_tot, cx_tot, cy_tot, n_pts = [], 0, 0, 0, 0
        for poly in polys:
            for ri, r in enumerate(poly):
                pts = [proj(p) for p in ring(r)]
                n_pts += len(pts)
                d_parts.append("M" + "L".join(f"{x:.1f},{y:.1f}" for x, y in pts) + "Z")
                if ri == 0:  # centroide del anillo exterior (fórmula del área)
                    a = cx = cy = 0
                    for (x1, y1), (x2, y2) in zip(pts, pts[1:] + pts[:1]):
                        c = x1 * y2 - x2 * y1
                        a += c
                        cx += (x1 + x2) * c
                        cy += (y1 + y2) * c
                    if a:
                        area_tot += a / 2
                        cx_tot += cx / 6
                        cy_tot += cy / 6
        p = g["properties"]
        feats.append({
            "cve": p["CVE_MUN"], "cvegeo": p["CVEGEO"], "nombre": p["NOMGEO"], "js": int(p["JS"]),
            "d": "".join(d_parts),
            "centroide": [round(cx_tot / area_tot, 1), round(cy_tot / area_tot, 1)] if area_tot else None,
            "vertices": n_pts,
        })
    feats.sort(key=lambda f: f["cve"])
    return {
        "viewBox": f"0 0 {width + 2 * pad:.0f} {height + 2 * pad:.0f}",
        "proyeccion": "Equirrectangular con corrección cos(latitud media) — sólo para despliegue",
        "simplificacion": f"Douglas-Peucker sobre arcos TopoJSON compartidos, tolerancia {tol}° (~{tol*111000:.0f} m)",
        "fuente": "Municipios_Queretaro_18 (TopoJSON/GeoJSON, CRS84) — proyecto QGIS/Power BI SNSP",
        "features": feats,
    }


def main():
    hoy = date.today().isoformat()
    filas = leer_base()
    audit = {"generado": hoy, "archivo": os.path.basename(RAW), "registros": len(filas)}

    # ---- Catálogo geográfico (GeoJSON + TopoJSON + Power BI) ----
    topo = json.load(open(TOPO, encoding="utf-8"))
    gj = json.load(open(GEOJSON, encoding="utf-8"))
    geo_props = {f["properties"]["CVE_MUN"]: f["properties"] for f in gj["features"]}
    topo_props = {g["properties"]["CVE_MUN"]: g["properties"] for g in topo["objects"]["municipios"]["geometries"]}
    wb = openpyxl.load_workbook(PBI, read_only=True, data_only=True)
    pbi_rows = list(wb["DIM_GEOGRAFIA_MUNICIPAL"].iter_rows(values_only=True))
    pbi_h = list(pbi_rows[0])
    pbi = {}
    for r in pbi_rows[1:]:
        d = dict(zip(pbi_h, r))
        pbi[f"{int(d['CVE_MUN']):03d}"] = d

    # ---- Validación de la base ----
    anios = sorted({int(f["Año"]) for f in filas})
    vistos = set()
    mun_info = {}
    nulos = 0
    for f in filas:
        if any(f[k] is None or f[k] == "" for k in f):
            nulos += 1
        clave = (f["Año"], f["Municipio"], f["Sexo"], f["Grupo de edad"])
        if clave in vistos:
            fail(f"Duplicado {clave}")
        vistos.add(clave)
        if f["Sexo"] not in SEXOS:
            fail(f"Sexo inesperado {f['Sexo']}")
        if f["Grupo de edad"] not in [g[0] for g in GRUPOS]:
            fail(f"Grupo inesperado {f['Grupo de edad']}")
        if not isinstance(f["Poblacion"], (int, float)) or f["Poblacion"] < 0:
            fail(f"Población inválida {f}")
        cve = str(f["Municipio"])[:3]
        js = int(str(f["Jurisdicción"])[:2])
        prev = mun_info.setdefault(cve, {"texto": f["Municipio"], "js": js, "js_texto": f["Jurisdicción"]})
        if prev["js"] != js:
            fail(f"Municipio {cve} en más de una jurisdicción")
    esperado = len(anios) * len(mun_info) * 2 * len(GRUPOS)
    audit["completitud"] = {"esperado": esperado, "encontrado": len(filas), "ok": esperado == len(filas)}
    audit["nulos"] = nulos
    audit["duplicados"] = len(filas) - len(vistos)
    if esperado != len(filas):
        fail("La base no está completa")

    # ---- Conciliación de claves ----
    concil = []
    for cve in sorted(set(mun_info) | set(geo_props) | set(topo_props) | set(pbi)):
        m, g, t, p = mun_info.get(cve), geo_props.get(cve), topo_props.get(cve), pbi.get(cve)
        ok = bool(m and g and t and p
                  and g["CVEGEO"] == "22" + cve and t["CVEGEO"] == "22" + cve
                  and str(p["CVEGEO_MUN"]) == "22" + cve
                  and int(g["JS"]) == m["js"] and int(t["JS"]) == m["js"]
                  and p["MUNICIPIO"] == g["NOMGEO"])
        concil.append({
            "cve": cve, "suive_texto": m and m["texto"], "suive_js": m and m["js"],
            "geo_nombre": g and g["NOMGEO"], "geo_cvegeo": g and g["CVEGEO"], "geo_js": g and int(g["JS"]),
            "topo_js": t and int(t["JS"]), "pbi_nombre": p and p["MUNICIPIO"],
            "pbi_municipio_suive": p and p["MUNICIPIO_SUIVE"], "coincide": ok,
            "texto_pbi_igual_a_suive": bool(m and p and p["MUNICIPIO_SUIVE"] == m["texto"]),
        })
        if not ok:
            fail(f"Clave {cve} no concilia entre fuentes")
    audit["conciliacion_claves"] = concil

    # ---- Decimales ----
    ruido = real = 0
    grupos_frac = defaultdict(int)
    for f in filas:
        v = float(f["Poblacion"])
        fr = abs(v - round(v))
        if 1e-9 < fr < 1e-4:
            ruido += 1
        elif fr >= 1e-4:
            real += 1
            grupos_frac[f["Grupo de edad"]] += 1
    par = defaultdict(float)
    for f in filas:
        if f["Grupo de edad"] in ("De 00 años", "De 01 a 04 años"):
            par[(f["Año"], f["Municipio"], f["Sexo"])] += float(f["Poblacion"])
    max_dev_par = max(abs(v - round(v)) for v in par.values())
    audit["decimales"] = {
        "filas_con_decimales": ruido + real,
        "ruido_punto_flotante": ruido,
        "fraccion_real": real,
        "grupos_con_fraccion_real": dict(grupos_frac),
        "max_desviacion_suma_00_mas_01_04_a_entero": max_dev_par,
        "interpretacion": "Ruido de exportación (±1e-6) en valores enteros; y prorrateo del grupo 0-4 entre "
                          "'De 00 años' y 'De 01 a 04 años' (su suma es entera). Se conservan los valores "
                          "originales y se redondea sólo al mostrar.",
    }

    # ---- Totales de referencia ----
    tot = defaultdict(float)
    tot_sexo = defaultdict(float)
    tot_js = defaultdict(float)
    for f in filas:
        tot[int(f["Año"])] += float(f["Poblacion"])
        tot_sexo[(int(f["Año"]), f["Sexo"])] += float(f["Poblacion"])
        tot_js[(int(f["Año"]), int(str(f["Jurisdicción"])[:2]))] += float(f["Poblacion"])
    audit["totales_anuales"] = {a: round(tot[a]) for a in anios}
    audit["totales_por_sexo"] = {a: {"Mujeres": round(tot_sexo[(a, "Femenino")]), "Hombres": round(tot_sexo[(a, "Masculino")])} for a in anios}
    audit["totales_por_jurisdiccion"] = {a: {NOMBRE_JS[j]: round(tot_js[(a, j)]) for j in sorted(NOMBRE_JS)} for a in anios}

    # ---- Conciliación con la extracción de Morbilidad (2022-2026) ----
    try:
        wbm = openpyxl.load_workbook(MORB, read_only=True, data_only=True)
        wsm = wbm[wbm.sheetnames[0]]
        mrows = list(wsm.iter_rows(values_only=True))
        mh = [str(h).strip() for h in mrows[0]]
        nuevo = defaultdict(float)
        for f in filas:
            nuevo[(int(f["Año"]), f["Municipio"], f["Grupo de edad"])] += float(f["Poblacion"])
        dif_max, n = 0.0, 0
        for r in mrows[1:]:
            d = dict(zip(mh, r))
            k = (int(d["Año"]), d["Municipio"], d["Grupo de edad"])
            if k in nuevo:
                n += 1
                dif_max = max(dif_max, abs(nuevo[k] - float(d["Población"])))
        audit["conciliacion_morbilidad_2022_2026"] = {"celdas_comparadas": n, "diferencia_maxima": dif_max, "identica": dif_max < 1e-6}
    except Exception as e:  # noqa
        audit["conciliacion_morbilidad_2022_2026"] = {"error": str(e)}

    # ---- Bundle compacto ----
    cves = sorted(mun_info)
    mun_idx = {c: i for i, c in enumerate(cves)}
    anio_idx = {a: i for i, a in enumerate(anios)}
    grp_idx = {g[0]: i for i, g in enumerate(GRUPOS)}
    registros = []
    for f in filas:
        registros.append([
            anio_idx[int(f["Año"])], mun_idx[str(f["Municipio"])[:3]], SEXOS[f["Sexo"]],
            grp_idx[f["Grupo de edad"]], float(f["Poblacion"]),
        ])
    registros.sort()

    data = {
        "meta": {
            "fuente": "SUIVE — base de población 2017-2026, estado de Querétaro (BASE_POBLACION_SUIVE_QRO_2017_2026_PLANA.xlsx, hoja POBLACION_BASE)",
            "fuente_corta": "Población SUIVE 2017-2026",
            "nota_fuente": "Base de población utilizada en el SUIVE. Su fuente metodológica (proyección) está pendiente de confirmar; no se atribuye a ninguna institución hasta entonces.",
            "generado": hoy,
            "registros": len(filas),
            "anios": anios,
            "nota_edad": "Se respetan los 11 grupos de edad de la fuente, de amplitud desigual (1, 4, 5, 5, 5, 5, 20, 5, 10, 5 años y 65+ abierto). No se generan quinquenales.",
            "nota_decimales": audit["decimales"]["interpretacion"],
            "nota_independencia": "Fuente distinta del padrón ASIS 2025 por localidad; ambas se presentan por separado y nunca se mezclan.",
        },
        "catalogos": {
            "anios": anios,
            "jurisdicciones": [{"id": j, "nombre": NOMBRE_JS[j], "etiqueta": f"JS {j} {NOMBRE_JS[j]}"} for j in sorted(NOMBRE_JS)],
            "municipios": [{"cve": c, "cvegeo": "22" + c, "nombre": geo_props[c]["NOMGEO"], "nombre_fuente": mun_info[c]["texto"], "js": mun_info[c]["js"]} for c in cves],
            "grupos_edad": [{"fuente": g[0], "etiqueta": g[1], "amplitud": g[2]} for g in GRUPOS],
            "sexos": ["Hombres", "Mujeres"],
        },
        "registros_campos": ["anio_idx", "municipio_idx", "sexo_idx(0=Hombres,1=Mujeres)", "grupo_idx", "poblacion_original"],
        "registros": registros,
        "auditoria": {
            "totales_anuales": audit["totales_anuales"],
            "conciliacion_morbilidad_2022_2026": audit["conciliacion_morbilidad_2022_2026"],
            "decimales": audit["decimales"],
        },
    }

    geo = construir_geo(topo)
    audit["geometria"] = {"features": len(geo["features"]), "vertices_totales": sum(f["vertices"] for f in geo["features"]),
                          "simplificacion": geo["simplificacion"], "viewBox": geo["viewBox"]}
    audit["errores"] = errores

    if errores:
        json.dump(audit, open(OUT_AUDIT, "w", encoding="utf-8"), ensure_ascii=False, indent=2)
        sys.exit(f"ABORTADO: {len(errores)} errores (ver {OUT_AUDIT})")

    hdr = ("/**\n * {path}\n * GENERADO AUTOMÁTICAMENTE por scripts/build_poblacion_historica.py — no editar a mano.\n"
           " * {desc}\n * Sin datos personales.\n */\n")
    with open(OUT_JS, "w", encoding="utf-8") as fh:
        fh.write(hdr.format(path="data/real/poblacion_historica_2017_2026.js",
                            desc="Fuente: BASE_POBLACION_SUIVE_QRO_2017_2026_PLANA.xlsx. Valores originales sin redondear."))
        fh.write("window.SNSP_POBLACION_HISTORICA_DATA = " + json.dumps(data, ensure_ascii=False, separators=(",", ":")) + ";\n")
    with open(OUT_GEO_JS, "w", encoding="utf-8") as fh:
        fh.write(hdr.format(path="data/real/geo_municipios_queretaro.js",
                            desc="Límites municipales de Querétaro (18) desde Municipios_Queretaro_18.topojson, simplificados para web."))
        fh.write("window.SNSP_GEO_QRO_MUNICIPIOS = " + json.dumps(geo, ensure_ascii=False, separators=(",", ":")) + ";\n")
    json.dump(data, open(OUT_JSON, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    json.dump(audit, open(OUT_AUDIT, "w", encoding="utf-8"), ensure_ascii=False, indent=2)
    print("OK", len(filas), "registros;", audit["totales_anuales"])
    print("Conciliación Morbilidad:", audit["conciliacion_morbilidad_2022_2026"])
    print("Geometría:", audit["geometria"])
    for f in geo["features"]:
        print(f["cve"], f["nombre"], f["vertices"], f["centroide"])


if __name__ == "__main__":
    main()
