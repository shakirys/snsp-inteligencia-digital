#!/usr/bin/env python3
"""
scripts/build_morbilidad.py
-----------------------------------------------------------------------
Transforma data/raw/CUBOS_DE_MORBILIDAD_2025_CON_CUBOS.xlsx (hoja
"Sheet 1") en:
  - data/processed/morbilidad_2025.json   (auditoría, legible)
  - data/real/morbilidad_2025.js          (window.SNSP_MORBILIDAD_DATA)

No modifica el archivo original. No calcula tasas ni incidencia (no hay
denominador poblacional validado). No inventa columnas: sólo usa año,
municipio, jurisdicción, institución, CLUES, unidad médica, mes,
epi-clave, padecimiento y casos confirmados, que son las que
efectivamente trae la base. Sexo y grupo de edad NO existen en esta
fuente y no se generan.

Complementa el nombre/estatus de la unidad médica mediante CLUES contra
ESTABLECIMIENTO_SALUD_202601.xlsx (relación 1 a 1 por CLUES, sin
duplicar casos).

La base sólo trae la entidad "Querétaro": este módulo es, por ahora,
exclusivamente estatal (ver README / metodología). La estructura queda
preparada para agregar entidad como dimensión el día que llegue una
base nacional, sin tener que rediseñar el cubo.
-----------------------------------------------------------------------
"""
import json
import re
import sys
import unicodedata
from collections import defaultdict

import openpyxl

RAW_MORBILIDAD = "data/raw/CUBOS_DE_MORBILIDAD_2025_CON_CUBOS.xlsx"
RAW_ESTABLECIMIENTOS = "data/raw/ESTABLECIMIENTO_SALUD_202601.xlsx"
OUT_JSON = "data/processed/morbilidad_2025.json"
OUT_JS = "data/real/morbilidad_2025.js"

CIE10_RE = re.compile(r"[A-Z]\d{2}(?:\.\d+)?")

# Accesos rápidos pedidos explícitamente. Sólo se activan si el código
# realmente aparece en el texto de "Padecimiento" de la base (nunca se
# inventa una relación). D05 se deja declarado pero se marcará como no
# disponible si no aparece.
ACCESOS_RAPIDOS_SOLICITADOS = ["C50", "C53", "D05", "D06", "N87"]

MESES_ORDEN = [
    "01 Enero", "02 Febrero", "03 Marzo", "04 Abril", "05 Mayo", "06 Junio",
    "07 Julio", "08 Agosto", "09 Septiembre", "10 Octubre", "11 Noviembre", "12 Diciembre",
]


def _strip_accents(s):
    return "".join(c for c in unicodedata.normalize("NFD", s) if unicodedata.category(c) != "Mn")


def _norm_key(s):
    return _strip_accents(s or "").upper().strip()


# Catálogo canónico de los 18 municipios de Querétaro (nombre correcto, con
# acentos y preposiciones en minúscula) — la base sólo trae MAYÚSCULAS sin
# acentos, así que se corrige aquí igual que build_data.py hace con
# "entidad" para CACU/Mama.
MUNICIPIOS_CANONICOS = {
    _norm_key(m): m for m in [
        "Amealco de Bonfil", "Arroyo Seco", "Cadereyta de Montes", "Colón",
        "Corregidora", "El Marqués", "Ezequiel Montes", "Huimilpan",
        "Jalpan de Serra", "Landa de Matamoros", "Pedro Escobedo",
        "Peñamiller", "Pinal de Amoles", "Querétaro", "San Joaquín",
        "San Juan del Río", "Tequisquiapan", "Tolimán",
    ]
}

# Las 4 jurisdicciones sanitarias que trae realmente la base (no se usa el
# catálogo de 3 jurisdicciones de config.js: es de otra fuente/demo).
JURISDICCIONES_CANONICAS = {
    _norm_key(j): j for j in [
        "Querétaro", "San Juan del Río", "Cadereyta de Montes", "Jalpan de Serra",
    ]
}

CONECTORES_ES = {"de", "del", "la", "las", "los", "y", "en", "para", "con", "el"}

# Las 8 instituciones que trae realmente la base (incluye "-" = no
# especificado). Se corrige sólo la acentuación; no se agrega ninguna que
# no esté en los datos.
INSTITUCIONES_CANONICAS = {
    _norm_key(i): i for i in [
        "Cruz Roja Mexicana",
        "Servicios Médicos Universitarios",
        "Instituto de Seguridad y Servicios Sociales para los Trabajadores del Estado",
        "Secretaría de Salud",
        "Sistema Nacional para el Desarrollo Integral de la Familia",
        "Servicios Médicos Privados",
        "Instituto Mexicano del Seguro Social",
    ]
}


def titlecase_institucion(s):
    if not s or s == "-":
        return s
    canon = INSTITUCIONES_CANONICAS.get(_norm_key(s))
    if canon:
        return canon
    words = s.strip().lower().split()
    out = []
    for i, w in enumerate(words):
        if i > 0 and w in CONECTORES_ES:
            out.append(w)
        else:
            out.append(w[:1].upper() + w[1:])
    return " ".join(out)


def canon_municipio(s):
    if not s or s == "-":
        return None
    return MUNICIPIOS_CANONICOS.get(_norm_key(s), titlecase_institucion(s))


def canon_jurisdiccion(s):
    if not s or s == "-":
        return None
    return JURISDICCIONES_CANONICAS.get(_norm_key(s), titlecase_institucion(s))


def main():
    print("Leyendo", RAW_MORBILIDAD, "...")
    wb = openpyxl.load_workbook(RAW_MORBILIDAD, read_only=True, data_only=True)
    ws = wb["Sheet 1"]

    padecimientos = {}  # epiclave -> padecimiento texto completo (tal cual la base)
    combos = defaultdict(int)  # (epiclave, anio, mun, juris, insti, clues, mes) -> casos
    clues_catalogo = {}  # clues -> {unidad_medica, municipio, jurisdiccion, institucion}
    anios = set()
    municipios = set()
    jurisdicciones = set()
    instituciones = set()
    total_casos = 0
    n_filas = 0

    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        (anio, entidad, clues, unidad, cve_tipo, tipologia, cve_insti, insti,
         cve_juris, juris, cve_mun, mun, cve_loc, loc, epiclave, idd, idfecha,
         mes, padec, quehacer, lat, lon, semana, casos) = row

        if not casos:
            continue
        casos = int(casos)
        n_filas += 1
        total_casos += casos

        mun = canon_municipio(mun)
        juris = canon_jurisdiccion(juris)
        insti = titlecase_institucion(insti)

        anios.add(int(anio))
        if mun and mun != "-":
            municipios.add(mun)
        if juris and juris != "-":
            jurisdicciones.add(juris)
        if insti and insti != "-":
            instituciones.add(insti)

        padecimientos.setdefault(str(epiclave), padec)

        key = (str(epiclave), int(anio), mun or "No especificado", juris or "No especificado",
               insti or "No especificado", clues or "No especificado", mes)
        combos[key] += casos

        if clues and clues not in clues_catalogo:
            clues_catalogo[clues] = {
                "clues": clues,
                "unidad_medica": unidad,
                "municipio": mun,
                "jurisdiccion": juris,
                "institucion": insti,
            }

    print(f"Filas con casos > 0: {n_filas} / total casos: {total_casos}")
    print(f"Padecimientos únicos: {len(padecimientos)}")
    print(f"Combinaciones únicas (cubo): {len(combos)}")
    print(f"CLUES únicos: {len(clues_catalogo)}")

    # ---- Catálogo de padecimientos + extracción de código(s) CIE-10 ----
    catalogo_padecimientos = []
    codigo_a_epiclaves = defaultdict(list)  # "C50" -> ["119"]  (para búsqueda y accesos rápidos)
    for epiclave, texto in padecimientos.items():
        m = re.search(r"\(([^)]*)\)\s*$", texto)
        cie10_texto = m.group(1) if m else ""
        nombre = texto[: texto.rfind("(")].strip() if m else texto
        codigos = sorted(set(CIE10_RE.findall(cie10_texto)))
        # También indexar por el prefijo de 3 caracteres (p.ej. "N87" a
        # partir de "N87.0"), para que un acceso rápido a "N87" encuentre
        # el padecimiento aunque el texto sólo traiga "N87.0-N87.1".
        prefijos = sorted(set(c[:3] for c in codigos))
        catalogo_padecimientos.append({
            "epiclave": epiclave,
            "padecimiento": nombre,
            "cie10_texto": cie10_texto,
            "cie10_codigos": codigos,
            "cie10_prefijos": prefijos,
        })
        for pref in prefijos:
            codigo_a_epiclaves[pref].append(epiclave)
        for full in codigos:
            codigo_a_epiclaves[full].append(epiclave)

    catalogo_padecimientos.sort(key=lambda r: r["padecimiento"])

    # ---- Accesos rápidos: sólo los que realmente están presentes ----
    accesos_rapidos = []
    for codigo in ACCESOS_RAPIDOS_SOLICITADOS:
        epiclaves = sorted(set(codigo_a_epiclaves.get(codigo, [])))
        accesos_rapidos.append({
            "codigo": codigo,
            "disponible": len(epiclaves) > 0,
            "epiclaves": epiclaves,
        })

    # ---- Enriquecimiento por CLUES desde ESTABLECIMIENTO_SALUD (Querétaro) ----
    print("Leyendo", RAW_ESTABLECIMIENTOS, "(sólo Querétaro) ...")
    wb2 = openpyxl.load_workbook(RAW_ESTABLECIMIENTOS, read_only=True, data_only=True)
    ws2 = wb2["CLUES_202601"]
    header = [c.value for c in next(ws2.iter_rows(min_row=1, max_row=1))]
    idx = {name: i for i, name in enumerate(header)}
    n_enriquecidos = 0
    for row in ws2.iter_rows(min_row=2, values_only=True):
        entidad = row[idx["ENTIDAD"]]
        if not entidad or "QUERETARO" not in str(entidad).upper():
            continue
        clues = row[idx["CLUES"]]
        if clues in clues_catalogo:
            clues_catalogo[clues]["nombre_unidad_oficial"] = row[idx["NOMBRE DE LA UNIDAD"]]
            clues_catalogo[clues]["tipo_establecimiento"] = row[idx["NOMBRE TIPO ESTABLECIMIENTO"]]
            clues_catalogo[clues]["nivel_atencion"] = row[idx["NIVEL ATENCION"]]
            clues_catalogo[clues]["estatus_operacion"] = row[idx["ESTATUS DE OPERACION"]]
            n_enriquecidos += 1
    print(f"CLUES enriquecidos con ESTABLECIMIENTO_SALUD: {n_enriquecidos} / {len(clues_catalogo)}")

    # ---- Cubo compacto (arreglos posicionales, no objetos, para tamaño) ----
    # [epiclave, anio, municipio, jurisdiccion, institucion, clues, mes, casos]
    cubo = [list(k) + [v] for k, v in combos.items()]

    data = {
        "meta": {
            "entidad": "Querétaro",
            "alcance": "estatal",
            "fuente": "CUBOS_DE_MORBILIDAD_2025_CON_CUBOS.xlsx — hoja 'Sheet 1'",
            "anios_disponibles": sorted(anios),
            "total_casos": total_casos,
            "total_padecimientos": len(catalogo_padecimientos),
            "total_clues": len(clues_catalogo),
            "nota": "Base exclusivamente de Querétaro. No incluye sexo ni grupo de "
                    "edad (no vienen en la fuente). No se calculan tasas ni "
                    "incidencia (sin denominador poblacional validado). "
                    "Estructura preparada para agregar 'entidad' como dimensión "
                    "cuando se reciba una base nacional, sin rediseñar el cubo.",
        },
        "catalogos": {
            "anios": sorted(anios),
            "municipios": sorted(municipios, key=lambda s: s),
            "jurisdicciones": sorted(jurisdicciones, key=lambda s: s),
            "instituciones": sorted(instituciones, key=lambda s: s),
            "meses": [m for m in MESES_ORDEN if m in {c[6] for c in combos}],
        },
        "catalogo_padecimientos": catalogo_padecimientos,
        "accesos_rapidos": accesos_rapidos,
        "clues_catalogo": clues_catalogo,
        "cubo_campos": ["epiclave", "anio", "municipio", "jurisdiccion", "institucion", "clues", "mes", "casos"],
        "cubo": cubo,
    }

    with open(OUT_JSON, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
    print("Escrito", OUT_JSON)

    with open(OUT_JS, "w", encoding="utf-8") as f:
        f.write(
            "/**\n"
            " * data/real/morbilidad_2025.js\n"
            " * GENERADO AUTOMÁTICAMENTE por scripts/build_morbilidad.py — no editar a mano.\n"
            " * Fuente: data/raw/CUBOS_DE_MORBILIDAD_2025_CON_CUBOS.xlsx (Querétaro).\n"
            " * No incluye datos personales, sólo conteos agregados.\n"
            " * -----------------------------------------------------------------------\n"
            " */\n"
            "window.SNSP_MORBILIDAD_DATA = "
        )
        json.dump(data, f, ensure_ascii=False, separators=(",", ":"))
        f.write(";\n")
    print("Escrito", OUT_JS)


if __name__ == "__main__":
    main()
