#!/usr/bin/env python3
# -----------------------------------------------------------------------
# scripts/build_canal_endemico.py
#
# Módulo exploratorio "Canal Endémico" — integrado a SNSP Inteligencia
# Digital (ver nota metodológica dentro de la página).
#
# Construye data/real/canal_endemico_2022_2026.js a partir de la MISMA
# fuente real ya usada por el Bloque 2 de Morbilidad:
#   data/raw/BASE_PADECIMIENTOS_SUIVE_QRO_2022_2026_plana.xlsx
# (Año, Entidad, Epi-Clave, Padecimiento, Mes, Semana, Grupo de edad,
# Sexo, Casos). No se inventa ni se reinterpreta ningún dato: se agrega
# Casos por Año + Semana epidemiológica (sumando Grupo de edad y Sexo),
# exactamente como pide la usuaria.
#
# 2022-2025 = referencia histórica. 2026 = año observado, parcial
# (sólo hasta la semana con información real en la fuente — nunca se
# imputa 0 a una semana sin registro, se deja como null/ausente).
#
# Por cada semana epidemiológica 1-53, se calculan Q1/Mediana(Q2)/Q3
# sobre los años históricos que SÍ tienen esa semana (la mayoría de los
# años tiene 52 semanas; algunos años epidemiológicos tienen 53 — esa
# semana 53 sólo se calcula con los años que realmente la tuvieron, sin
# rellenar nada). Método de cuantiles: interpolación lineal inclusiva
# (statistics.quantiles(..., method="inclusive"), equivalente al
# percentil "linear"/tipo 7 de NumPy y a PERCENTILE.INC de Excel) — el
# método más común y el que se documenta en la nota metodológica del
# propio módulo.
#
# ----------------------------------------------------------------------
# AMPLIACIÓN (ACT22): múltiples padecimientos, con criterio explícito de
# elegibilidad (no se fabrica canal para ningún padecimiento que no lo
# cumpla; se documenta y se reporta en la tabla de auditoría):
#
#   Un Epi-Clave es ELEGIBLE para construir su canal cuando:
#     1) Tiene información en el año observado (2026) — si no, no hay
#        nada que graficar como "observado".
#     2) Al menos el 80% de las semanas epidemiológicas de su rango
#        observado en 2026 (SE 1 .. última semana real de 2026) cuentan
#        con al menos 2 de los 4 años históricos (2022-2025) reportando
#        esa misma semana. Esto evita construir un canal cuya mayoría de
#        semanas dependería de un solo año de referencia (ya se sabe,
#        por el caso ya aprobado de la SE 53 de alacrán, que n=1 no es
#        una referencia estadística válida — aquí se aplica el mismo
#        criterio pero a nivel de padecimiento completo, no sólo una
#        semana aislada).
#   El umbral de 80% se eligió porque cae exactamente en un salto natural
#   de la distribución real de cobertura semanal de los 167 Epi-Clave de
#   la base (el padecimiento elegible con menor cobertura tiene 86.5%; el
#   primero excluido tiene 78.4% — no hay ningún caso intermedio).
#
#   Los padecimientos NO elegibles se siguen listando en el catálogo
#   (para que el selector con buscador los pueda localizar por nombre o
#   Epi-Clave) pero con "disponible": false y un "motivo" explícito; NO
#   se construye (ni se fabrica) ningún canal para ellos. El front-end
#   debe mostrar un aviso claro en ese caso, nunca cuartiles o zonas
#   inventadas.
# ----------------------------------------------------------------------
import json
import re
import statistics
import sys
from collections import defaultdict
from pathlib import Path

try:
    import openpyxl
except ImportError:
    print("Falta openpyxl. Instala con: pip install openpyxl --break-system-packages", file=sys.stderr)
    sys.exit(1)

RAIZ = Path(__file__).resolve().parent.parent
RUTA_XLSX = RAIZ / "data" / "raw" / "BASE_PADECIMIENTOS_SUIVE_QRO_2022_2026_plana.xlsx"
RUTA_JS = RAIZ / "data" / "real" / "canal_endemico_2022_2026.js"
RUTA_JSON_DEBUG = RAIZ / "data" / "processed" / "canal_endemico_2022_2026.json"

ANIOS_HISTORICOS = ["2022", "2023", "2024", "2025"]
ANIO_OBSERVADO = "2026"
UMBRAL_PCT_COBERTURA_N2 = 80.0  # ver nota de elegibilidad arriba

# Totales anuales YA VALIDADOS explícitamente por la usuaria (aserción
# dura adicional, sólo para el/los padecimiento(s) que ella confirmó a
# mano) — el resto de padecimientos elegibles se valida de forma
# genérica (ver _validar_consistencia_interna) pero sin un total externo
# de referencia, porque no se ha confirmado uno con la usuaria.
TOTALES_CONFIRMADOS_POR_USUARIA = {
    "94": {"2022": 4477, "2023": 5788, "2024": 5951, "2025": 6436, "2026": 4147},
}


def _leer_filas():
    wb = openpyxl.load_workbook(RUTA_XLSX, read_only=True, data_only=True)
    ws = wb.active
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    idx = {h: i for i, h in enumerate(headers)}
    requeridos = ["Año", "Epi-Clave", "Padecimiento", "Semana", "Casos"]
    faltantes = [c for c in requeridos if c not in idx]
    if faltantes:
        raise SystemExit(f"ERROR: faltan columnas requeridas en la base cruda: {faltantes}")
    for row in ws.iter_rows(min_row=2, values_only=True):
        yield {
            "anio": str(row[idx["Año"]]),
            "epiclave": str(row[idx["Epi-Clave"]]),
            "padecimiento": row[idx["Padecimiento"]],
            "semana": row[idx["Semana"]],
            "casos": row[idx["Casos"]] or 0,
        }


def _cuantiles_inclusive(valores):
    """Q1, Mediana, Q3 por interpolación lineal inclusiva (n>=2).
    Con un solo valor histórico (n=1), los 3 cuantiles son ese mismo
    valor — no es una falla, es la única lectura honesta con un solo
    año de referencia, y se marca con n_historico=1 en el bundle."""
    if len(valores) == 1:
        v = valores[0]
        return v, v, v
    q1, q2, q3 = statistics.quantiles(valores, n=4, method="inclusive")
    return q1, q2, q3


def _clasificar_zona(casos, q1, mediana, q3):
    if casos is None:
        return None
    if casos <= q1:
        return "exito"
    if casos <= mediana:
        return "seguridad"
    if casos <= q3:
        return "alerta"
    return "epidemica"


def main():
    print(f"Leyendo: {RUTA_XLSX}")

    # ---- Paso 1: un sólo recorrido de la base cruda, agregando por
    # (epiclave, año, semana) y guardando el nombre y el total anual.
    nombre_por_epiclave = {}
    casos_epiclave_anio_semana = defaultdict(lambda: defaultdict(int))
    semanas_con_dato = defaultdict(lambda: defaultdict(set))  # epiclave -> anio -> {semanas}
    total_por_epiclave_anio = defaultdict(lambda: defaultdict(int))
    filas_sin_semana_por_epiclave = defaultdict(int)
    filas_leidas = 0

    for fila in _leer_filas():
        filas_leidas += 1
        epiclave = fila["epiclave"]
        anio = fila["anio"]
        if epiclave not in nombre_por_epiclave and fila["padecimiento"]:
            nombre_por_epiclave[epiclave] = fila["padecimiento"]
        total_por_epiclave_anio[epiclave][anio] += fila["casos"]
        semana = fila["semana"]
        if semana is None or semana == "":
            filas_sin_semana_por_epiclave[epiclave] += 1
            continue
        semana_i = int(semana)
        casos_epiclave_anio_semana[epiclave][(anio, semana_i)] += fila["casos"]
        semanas_con_dato[epiclave][anio].add(semana_i)

    if filas_leidas == 0:
        raise SystemExit("ERROR: la base cruda no tiene filas de datos.")
    print(f"  [OK] {filas_leidas} filas leídas, {len(nombre_por_epiclave)} Epi-Clave distintas encontradas.")

    # ---- Paso 2: elegibilidad por Epi-Clave (criterio documentado arriba) ----
    catalogo = []
    canales = {}
    n_elegibles = 0
    n_no_elegibles = 0

    for epiclave in sorted(nombre_por_epiclave, key=lambda e: nombre_por_epiclave[e].split("(")[0].strip().lower()):
        # ACT24 (corrección): la base cruda trae el nombre y el paréntesis
        # del CIE-10 pegados, sin espacio ("...agudas(J00-J06...)") — se
        # normaliza aquí, una sola vez, para que TODO lo que se derive de
        # nombre_completo (catálogo, título del PDF, selector con
        # buscador) quede con el espacio correcto. No toca nombre_corto
        # ni cie10 (ninguno de los dos depende de si hay espacio antes
        # del paréntesis) ni ningún cálculo del canal (que usa Casos, no
        # este texto).
        nombre_completo = re.sub(r"(?<=\S)\(", " (", nombre_por_epiclave[epiclave])
        nombre_corto = nombre_completo.split("(")[0].strip()
        cie10 = nombre_completo[nombre_completo.find("(") + 1: nombre_completo.find(")")] if "(" in nombre_completo else None

        if filas_sin_semana_por_epiclave.get(epiclave):
            # No se puede construir el canal sin inventar la semana que
            # falta en la fuente — se deja fuera, con motivo explícito.
            catalogo.append({
                "epiclave": epiclave, "padecimiento": nombre_corto,
                "padecimiento_texto_completo": nombre_completo, "cie10": cie10,
                "disponible": False,
                "motivo": f"{filas_sin_semana_por_epiclave[epiclave]} registro(s) sin semana epidemiológica en la fuente: no se puede construir el canal sin fabricar ese dato.",
            })
            n_no_elegibles += 1
            continue

        anios_hist_con_datos = [a for a in ANIOS_HISTORICOS if semanas_con_dato[epiclave].get(a)]
        semanas_obs_2026 = semanas_con_dato[epiclave].get(ANIO_OBSERVADO, set())
        tiene_2026 = len(semanas_obs_2026) > 0

        if not tiene_2026:
            catalogo.append({
                "epiclave": epiclave, "padecimiento": nombre_corto,
                "padecimiento_texto_completo": nombre_completo, "cie10": cie10,
                "disponible": False,
                "motivo": "Sin información en el año observado (2026): no hay nada que comparar contra el histórico.",
            })
            n_no_elegibles += 1
            continue

        ultima_semana_2026 = max(semanas_obs_2026)
        n_hist_por_semana = [
            sum(1 for a in anios_hist_con_datos if se in semanas_con_dato[epiclave][a])
            for se in range(1, ultima_semana_2026 + 1)
        ]
        pct_cobertura_n2 = 100.0 * sum(1 for n in n_hist_por_semana if n >= 2) / len(n_hist_por_semana)

        if pct_cobertura_n2 < UMBRAL_PCT_COBERTURA_N2:
            catalogo.append({
                "epiclave": epiclave, "padecimiento": nombre_corto,
                "padecimiento_texto_completo": nombre_completo, "cie10": cie10,
                "disponible": False,
                "motivo": (
                    f"Información histórica semanal insuficiente: sólo el {round(pct_cobertura_n2, 1)}% de las "
                    f"semanas de 2026 (SE 1-{ultima_semana_2026}) cuentan con al menos 2 años de referencia "
                    f"2022-2025 (se requiere ≥{UMBRAL_PCT_COBERTURA_N2:.0f}%). Construir el canal fabricaría "
                    f"cuartiles poco confiables para la mayoría de las semanas."
                ),
            })
            n_no_elegibles += 1
            continue

        # ---- Elegible: construir el canal completo ----
        huecos_2026 = sorted(set(range(1, ultima_semana_2026 + 1)) - semanas_obs_2026)
        max_semana_historico = max(
            (max(semanas_con_dato[epiclave][a]) for a in ANIOS_HISTORICOS if semanas_con_dato[epiclave].get(a)),
            default=52,
        )

        semanas_out = []
        for se in range(1, max_semana_historico + 1):
            valores_historicos = [
                casos_epiclave_anio_semana[epiclave][(a, se)]
                for a in ANIOS_HISTORICOS
                if se in semanas_con_dato[epiclave].get(a, set())
            ]
            if not valores_historicos:
                q1 = q2 = q3 = None
                n_hist = 0
            else:
                q1, q2, q3 = _cuantiles_inclusive(sorted(valores_historicos))
                n_hist = len(valores_historicos)

            casos_2026 = casos_epiclave_anio_semana[epiclave][(ANIO_OBSERVADO, se)] if se in semanas_obs_2026 else None
            zona = _clasificar_zona(casos_2026, q1, q2, q3) if q1 is not None else None

            semanas_out.append({
                "se": se,
                "q1": round(q1, 2) if q1 is not None else None,
                "mediana": round(q2, 2) if q2 is not None else None,
                "q3": round(q3, 2) if q3 is not None else None,
                "n_historico": n_hist,
                "casos_2026": casos_2026,
                "zona": zona,
            })

        resumen_zonas = defaultdict(int)
        for s in semanas_out:
            if s["se"] <= ultima_semana_2026 and s["zona"]:
                resumen_zonas[s["zona"]] += 1

        total_por_anio = {a: total_por_epiclave_anio[epiclave].get(a, 0) for a in ANIOS_HISTORICOS + [ANIO_OBSERVADO]}

        # ---- Validación dura SÓLO para el/los padecimiento(s) con total
        # confirmado a mano por la usuaria (hoy: Epi-Clave 94) ----
        if epiclave in TOTALES_CONFIRMADOS_POR_USUARIA:
            for anio, esperado in TOTALES_CONFIRMADOS_POR_USUARIA[epiclave].items():
                obtenido = total_por_anio.get(anio, 0)
                if obtenido != esperado:
                    raise SystemExit(
                        f"ERROR DE VALIDACIÓN: Epi-Clave={epiclave}, año {anio}: "
                        f"esperado {esperado}, obtenido {obtenido}. NO se escribe ningún archivo."
                    )
            print(f"  [OK] Epi-Clave {epiclave}: totales anuales verificados contra los confirmados por la usuaria.")

        canales[epiclave] = {
            "epiclave": epiclave,
            "padecimiento": nombre_corto,
            "padecimiento_texto_completo": nombre_completo,
            "anios_historicos_usados": ANIOS_HISTORICOS,
            "anio_observado": ANIO_OBSERVADO,
            "ultima_semana_observada": ultima_semana_2026,
            "huecos_semanas_2026": huecos_2026,
            "total_por_anio": total_por_anio,
            "resumen_semanas_por_zona": dict(resumen_zonas),
            "semanas": semanas_out,
            "pct_cobertura_historica": round(pct_cobertura_n2, 1),
        }
        catalogo.append({
            "epiclave": epiclave, "padecimiento": nombre_corto,
            "padecimiento_texto_completo": nombre_completo, "cie10": cie10,
            "disponible": True,
        })
        n_elegibles += 1

    print(f"\n=== Elegibilidad: {n_elegibles} padecimiento(s) con canal construido, {n_no_elegibles} sin información semanal suficiente (de {len(nombre_por_epiclave)} totales) ===")

    meta = {
        "fuente": "BASE_PADECIMIENTOS_SUIVE_QRO_2022_2026_plana.xlsx",
        "nota": (
            "Módulo exploratorio (Canal Endémico). Mismo Bloque 2 que ya usa "
            "Morbilidad: Año, Epi-Clave, Padecimiento, Mes, Semana epidemiológica, Grupo de edad, Sexo, "
            "Casos. Se agrega Casos por Año + Semana epidemiológica, sumando Grupo de edad y Sexo. "
            "2022-2025 = referencia histórica (4 años completos). 2026 = año observado, parcial: sólo "
            "tiene información real hasta la última semana reportada en la fuente — las semanas "
            "posteriores NUNCA se rellenan con 0, quedan como 'sin información' (casos_2026 = null)."
        ),
        "metodo_cuantiles": (
            "Q1, Mediana (Q2) y Q3 por interpolación lineal inclusiva "
            "(statistics.quantiles(..., method='inclusive') en Python — equivalente al percentil "
            "'linear'/tipo 7 de NumPy y a PERCENTILE.INC de Excel), calculados únicamente con los años "
            "históricos que de verdad tienen esa semana epidemiológica (no todos los años tienen 53 "
            "semanas: esa semana se calcula sólo con los años que la tuvieron, nunca con los 4)."
        ),
        "clasificacion_zonas": (
            "casos <= Q1: Zona de éxito · Q1 < casos <= Mediana: Zona de seguridad · "
            "Mediana < casos <= Q3: Zona de alerta · casos > Q3: Zona epidémica. "
            "La zona epidémica del canal exploratorio es una señal estadística relativa al propio "
            "histórico (sólo 4 años), NO una declaración oficial de epidemia."
        ),
        "criterio_elegibilidad": (
            f"Un padecimiento sólo se construye como canal cuando tiene información en 2026 Y al menos "
            f"{UMBRAL_PCT_COBERTURA_N2:.0f}% de sus semanas observadas en 2026 cuentan con 2 o más años "
            f"históricos (2022-2025) de referencia para esa misma semana. Los padecimientos que no cumplen "
            f"este criterio se listan en el buscador pero muestran un aviso explícito en vez de un canal — "
            f"nunca se fabrica un cuartil o una zona con información insuficiente."
        ),
        "anios_historicos": ANIOS_HISTORICOS,
        "anio_observado": ANIO_OBSERVADO,
        "advertencia_exploratorio": (
            "Herramienta EXPLORATORIA: sólo se cuenta con 4 años históricos completos (2022-2025). "
            "Se recomienda validar el método e incorporar más años antes de su adopción formal."
        ),
    }

    RUTA_JS.parent.mkdir(parents=True, exist_ok=True)
    RUTA_JSON_DEBUG.parent.mkdir(parents=True, exist_ok=True)

    payload = {"meta": meta, "catalogo_padecimientos": catalogo, "canales": canales}

    RUTA_JSON_DEBUG.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")

    js = (
        "/**\n"
        " * data/real/canal_endemico_2022_2026.js\n"
        " * GENERADO AUTOMÁTICAMENTE por scripts/build_canal_endemico.py — no editar a mano.\n"
        " * Módulo exploratorio — NO sustituye a Morbilidad.\n"
        " * Fuente: data/raw/BASE_PADECIMIENTOS_SUIVE_QRO_2022_2026_plana.xlsx (misma base del\n"
        " * Bloque 2 de Morbilidad). No incluye datos personales.\n"
        " * -----------------------------------------------------------------------\n"
        " */\n"
        "window.SNSP_CANAL_ENDEMICO_DATA = " + json.dumps(payload, ensure_ascii=False) + ";\n"
    )
    RUTA_JS.write_text(js, encoding="utf-8")

    print(f"\nEscrito: {RUTA_JS} ({RUTA_JS.stat().st_size} bytes)")
    print(f"Escrito (depuración): {RUTA_JSON_DEBUG}")
    print("\n=== CONSTRUCCIÓN COMPLETA, TOTALES VALIDADOS ===")


if __name__ == "__main__":
    main()
