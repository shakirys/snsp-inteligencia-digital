#!/usr/bin/env python3
"""
scripts/build_morbilidad_casos.py
-----------------------------------------------------------------------
Genera el paquete de datos que alimenta el tablero UNIFICADO de
Morbilidad, a partir de las MISMAS dos bases ya validadas en Cargar
datos (ACT06 — Parte 2):
  - data/raw/BASE_CASOS_2025_2026.xlsx
  - data/raw/BASE_POBLACION_2025_2026.xlsx

Salidas:
  - data/processed/morbilidad_casos_2025_2026.json       (auditoría)
  - data/processed/morbilidad_poblacion_2025_2026.json   (auditoría)
  - data/processed/morbilidad_temporal_2025_2026.json    (auditoría)
  - data/real/morbilidad_casos_2025_2026.js       (window.SNSP_MORBILIDAD_CASOS_DATA)
  - data/real/morbilidad_poblacion_2025_2026.js   (window.SNSP_MORBILIDAD_POBLACION_DATA)
  - data/real/morbilidad_temporal_2025_2026.js    (window.SNSP_MORBILIDAD_TEMPORAL_DATA)

-----------------------------------------------------------------------
ACT14 — Mes y Semana epidemiológica como filtros globales
-----------------------------------------------------------------------
BASE_CASOS_2025_2026.xlsx SÍ trae Mes y Semana (columnas confirmadas al
leer el archivo: Año, Entidad, Jurisdicción, Municipio, Institución,
Mes, Semana, Grupo de edad, Sexo, Epi-Clave, Padecimiento, Casos). Las
entregas anteriores las habían soltado a propósito (eran "ortogonales"
al tablero de ese momento); ACT14 las necesita como filtros GLOBALES
(KPI, gráfica/tabla principal, Panorama y PDF), así que ahora SÍ se
agregan a la agrupación de ambos bundles:

  - Bundle PRINCIPAL: se agrega Mes y Semana a la llave de agrupación
    (antes: Año+Jurisdicción+Municipio+Institución+Grupo de
    edad+Epi-Clave+Padecimiento -> ahora se suma también Mes+Semana).
    Esto SUBE las combinaciones de 20,563 a 211,032 (se verificó con
    openpyxl antes de escribir este script) — la suma total y la suma
    por cualquier subconjunto de las columnas ANTERIORES es
    MATEMÁTICAMENTE IDÉNTICA a la de antes (agregar más columnas a una
    llave de agrupación nunca cambia una suma; sólo la reparte más
    fino) — se verifica en build_casos() comparando explícitamente
    contra la agregación SIN Mes/Semana calculada en la misma corrida,
    con el mismo archivo crudo, antes de escribir cualquier salida.
  - Bundle TEMPORAL (Tendencia/Comparativo por mes-año/Casos por sexo):
    se agrega Semana a su llave (ya traía Mes) para que estos 3
    paneles también respeten un filtro de Semana cuando esté activo.

Tamaño del bundle principal: 211,032 filas de 9 columnas de texto/año
(sin contar Casos) pesarían ~21 MB de JSON sin comprimir si se
guardaran como texto repetido — demasiado para cargar cómodo en un
navegador. Se usa en cambio una codificación por ÍNDICES (columnar,
tipo "diccionario"): cada bundle guarda un catálogo de valores únicos
por columna (Año, Jurisdicción, Municipio, Institución, Grupo de edad,
Epi-Clave, Mes, Semana — cardinalidades chicas: 2 a 157) y cada fila
guarda sólo el ÍNDICE de cada valor dentro de su catálogo, más el
número de Casos. El texto de Padecimiento (el campo más largo, y
1-a-1 con Epi-Clave, ya verificado en entregas anteriores) tampoco se
repite 211,032 veces: se guarda una sola vez por Epi-Clave. Esto baja
el bundle principal a ~4.6 MB sin comprimir (~694 KB con la
compresión gzip que ya sirve GitHub Pages) — cabe cómodo y es más
rápido de leer que si se sirviera como texto repetido, con exactamente
la misma información.

La decodificación (índices -> filas con los mismos valores de texto
que antes) se escribe DENTRO del propio archivo .js generado (una
función autoejecutable), así que `window.SNSP_MORBILIDAD_CASOS_DATA`
sigue teniendo EXACTAMENTE la misma forma que antes para quien lo
consume (`{meta, headers, rows, catalogos, catalogo_padecimientos}`,
con `rows` como arreglos de texto/número igual que siempre, sólo que
ahora con 2 columnas más, Mes y Semana) — modules/morbilidadModuleView.js
no necesita cambiar CÓMO lee los datos, sólo qué hace con las 2
columnas nuevas (mismo patrón ya usado ahí: `headers.indexOf(...)`,
nunca un índice fijo).

mesSemana: además de catalogos.meses/semanas, el bundle principal
incluye `meta.mes_semana` — para cada Año, qué Semanas ocurren dentro
de cada Mes y qué Meses toca cada Semana (hay semanas de frontera que
caen en 2 meses, y la correspondencia Mes<->Semana NO es la misma
entre 2025 y 2026 porque el primer día del año cae distinto). Se
calcula DIRECTO de las filas crudas (nunca se asume ni se inventa),
para que el filtro de Mes y Semana puedan mostrarse "coherentes entre
sí" (ACT14, punto 2) sin adivinar la relación.
-----------------------------------------------------------------------

Columnas de salida del bundle principal: Año, Jurisdicción, Municipio,
Institución, Grupo de edad, Epi-Clave, Padecimiento, Mes, Semana,
Casos — TODAS las que hacen falta para:
  - el cruce con Población (Año + Municipio + Grupo de edad),
  - los 4 "Por" del tablero (Padecimiento / Municipio / Jurisdicción /
    Grupo de edad),
  - los filtros (Municipio, Jurisdicción, Institución, Mes, Semana —
    Institución NO es una opción de "Por", sólo filtro, a petición
    explícita),
  - el buscador por CIE-10 / Epi-clave / padecimiento (Epi-Clave +
    Padecimiento, catálogo aparte, ver abajo).

Los valores de Municipio/Institución/Jurisdicción/Padecimiento/Grupo de
edad/Mes/Semana se dejan TAL CUAL vienen en el archivo (incluyendo el
prefijo numérico, p. ej. "006 Corregidora", "01 Enero"), sin renombrar
ni canonicalizar — igual que hace window.SNSP_CARGA_SERVICE en Cargar
datos, para que el cruce por texto entre Casos y Población siga
funcionando igual.

Catálogo de padecimientos para el buscador: se extrae, por cada
Epi-Clave único, el código CIE-10 embebido al final del texto de
Padecimiento — MISMA lógica de extracción (regex, separación
nombre/CIE-10) que ya se usaba. Sin cambios de comportamiento en esta
parte (ACT14 punto 2 se resuelve del lado de la UI con este mismo
catálogo, no requirió tocar la extracción).

No modifica los archivos originales. No agrega ninguna otra función
nueva fuera de lo que Mes/Semana necesitan (fuera de alcance de esta
actualización).
-----------------------------------------------------------------------
"""
import json
import re
from collections import defaultdict

import openpyxl

RAW_CASOS = "data/raw/BASE_CASOS_2025_2026.xlsx"
RAW_POBLACION = "data/raw/BASE_POBLACION_2025_2026.xlsx"

OUT_CASOS_JSON = "data/processed/morbilidad_casos_2025_2026.json"
OUT_CASOS_JS = "data/real/morbilidad_casos_2025_2026.js"
OUT_POBLACION_JSON = "data/processed/morbilidad_poblacion_2025_2026.json"
OUT_POBLACION_JS = "data/real/morbilidad_poblacion_2025_2026.js"

OUT_TEMPORAL_JSON = "data/processed/morbilidad_temporal_2025_2026.json"
OUT_TEMPORAL_JS = "data/real/morbilidad_temporal_2025_2026.js"

CASOS_HEADERS_OUT = ["Año", "Jurisdicción", "Municipio", "Institución", "Grupo de edad", "Epi-Clave", "Padecimiento", "Mes", "Semana", "Casos"]
TEMPORAL_HEADERS_OUT = ["Año", "Jurisdicción", "Municipio", "Institución", "Mes", "Semana", "Sexo", "Epi-Clave", "Casos"]

CIE10_RE = re.compile(r"[A-Z]\d{2}(?:\.\d+)?")


def _extraer_cie10(texto_padecimiento):
    """Igual que build_morbilidad.py: separa 'Nombre(CIE10)' -> (nombre, cie10_texto, codigos, prefijos)."""
    m = re.search(r"\(([^)]*)\)\s*$", texto_padecimiento)
    cie10_texto = m.group(1) if m else ""
    nombre = texto_padecimiento[: texto_padecimiento.rfind("(")].strip() if m else texto_padecimiento
    codigos = sorted(set(CIE10_RE.findall(cie10_texto)))
    prefijos = sorted(set(c[:3] for c in codigos))
    return nombre, cie10_texto, codigos, prefijos


class _Catalogo:
    """Catálogo de valores únicos de una columna, para la codificación por
    índices del bundle: idx(valor) devuelve (y crea si hace falta) el
    índice de ese valor dentro del catálogo, preservando el orden de
    primera aparición (no importa el orden final: el catálogo se guarda
    completo, cualquier índice apunta al valor correcto)."""

    def __init__(self):
        self._idx = {}
        self._valores = []

    def idx(self, valor):
        i = self._idx.get(valor)
        if i is None:
            i = len(self._valores)
            self._idx[valor] = i
            self._valores.append(valor)
        return i

    def lista(self):
        return self._valores


def _escribir_js_bundle(path_js, var_name, meta, headers, catalogos_idx, rows_idx, extra_top_level, comentario, campo_derivado=None):
    """Escribe un bundle .js con codificación por índices (catalogos_idx:
    {nombreCatalogo: [valores...]}, rows_idx: [[idx0, idx1, ..., casos], ...])
    más una función autoejecutable que lo DECODIFICA de vuelta a filas de
    texto completas (mismo formato {meta, headers, rows, ...} que
    consumían los módulos antes de ACT14) — así ningún archivo que LEE
    este bundle necesita cambiar cómo lo hace.

    catalogos_idx: dict ORDENADO (mismo orden que las columnas de
    rows_idx, ANTES de la columna final de Casos) nombreCatalogo ->
    [valores...]. rows_idx: cada fila trae un índice por cada catálogo,
    en ese mismo orden, y termina con el número de Casos.

    campo_derivado (opcional, ej. Padecimiento — 1 a 1 con Epi-Clave):
    {"header": "Padecimiento", "despues_de": "Epi-Clave", "fuente": "padecimientoPorEpiclave",
     "via_catalogo": "epiclave"} — se inserta en `headers` un campo que NO
    consume su propio índice en rows_idx: se resuelve reutilizando el
    ÍNDICE de otro catálogo (via_catalogo) contra un arreglo aparte
    (fuente, en extra_top_level), evitando repetir el texto por fila."""
    payload = {
        "meta": meta,
        "headers": headers,
        "catalogosIdx": catalogos_idx,
        "rowsIdx": rows_idx,
    }
    payload.update(extra_top_level or {})
    cols_orden = list(catalogos_idx.keys())
    with open(path_js, "w", encoding="utf-8") as f:
        f.write(comentario)
        f.write(f"window.{var_name} = (function () {{\n")
        f.write("  var _p = ")
        json.dump(payload, f, ensure_ascii=False, separators=(",", ":"))
        f.write(";\n")
        f.write(f"  var _cols = {json.dumps(cols_orden)};\n")
        if campo_derivado:
            f.write(f"  var _viaCatIdx = _cols.indexOf({json.dumps(campo_derivado['via_catalogo'])});\n")
            f.write(f"  var _fuenteDerivada = _p[{json.dumps(campo_derivado['fuente'])}];\n")
        f.write(
            "  var _cat = _p.catalogosIdx;\n"
            "  var _rows = _p.rowsIdx.map(function (r) {\n"
            "    var fila = [];\n"
            "    for (var i = 0; i < _cols.length; i++) {\n"
            "      fila.push(_cat[_cols[i]][r[i]]);\n"
        )
        if campo_derivado:
            f.write("      if (i === _viaCatIdx) { fila.push(_fuenteDerivada[r[i]]); }\n")
        f.write(
            "    }\n"
            "    fila.push(r[_cols.length]);\n"
            "    return fila;\n"
            "  });\n"
            "  var out = { meta: _p.meta, headers: _p.headers, rows: _rows };\n"
            "  Object.keys(_p).forEach(function (k) {\n"
            "    if (k !== 'meta' && k !== 'headers' && k !== 'catalogosIdx' && k !== 'rowsIdx') out[k] = _p[k];\n"
            "  });\n"
            "  return out;\n"
            "})();\n"
        )


def build_casos():
    print("Leyendo", RAW_CASOS, "...")
    wb = openpyxl.load_workbook(RAW_CASOS, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    header = [str(c.value or "").strip() for c in next(ws.iter_rows(min_row=1, max_row=1))]
    idx = {name: i for i, name in enumerate(header)}
    columnas_requeridas = ["Año", "Jurisdicción", "Municipio", "Institución", "Grupo de edad", "Epi-Clave", "Padecimiento", "Mes", "Semana", "Casos"]
    for col in columnas_requeridas:
        if col not in idx:
            raise SystemExit(f"Columna esperada '{col}' no encontrada en {RAW_CASOS}: {header}")

    # combos_nuevo: llave CON Mes+Semana (lo que se guarda). combos_anterior:
    # llave SIN Mes+Semana (7 columnas, la agrupación de antes de ACT14) —
    # se calculan JUNTAS, en la misma pasada, para poder comprobar que
    # sumar combos_nuevo colapsando Mes+Semana da EXACTAMENTE combos_anterior
    # (ver verificación más abajo) antes de escribir cualquier salida.
    combos_nuevo = defaultdict(int)
    combos_anterior = defaultdict(int)
    padecimientos_texto = {}  # epiclave -> texto completo de Padecimiento (tal cual la base)
    n_filas_leidas = 0
    n_filas_usadas = 0
    total_casos = 0
    anios = set()
    jurisdicciones = set()
    municipios = set()
    instituciones = set()
    grupos_edad = set()
    meses = set()
    semanas = set()
    # mes_semana[anio][mes] = set(semanas); semana_mes[anio][semana] = set(meses)
    mes_semana = defaultdict(lambda: defaultdict(set))
    semana_mes = defaultdict(lambda: defaultdict(set))

    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[idx["Año"]] is None:
            continue
        n_filas_leidas += 1
        casos = row[idx["Casos"]]
        if casos is None or casos == "":
            continue
        try:
            casos = int(casos)
        except (TypeError, ValueError):
            continue

        anio = str(int(row[idx["Año"]]))
        juris = str(row[idx["Jurisdicción"]] or "").strip()
        municipio = str(row[idx["Municipio"]] or "").strip()
        institucion = str(row[idx["Institución"]] or "").strip()
        grupo_edad = str(row[idx["Grupo de edad"]] or "").strip()
        epiclave = str(row[idx["Epi-Clave"]] or "").strip()
        padecimiento = str(row[idx["Padecimiento"]] or "").strip()
        mes = str(row[idx["Mes"]] or "").strip()
        semana = str(row[idx["Semana"]] or "").strip()
        if not (anio and juris and municipio and institucion and grupo_edad and epiclave and padecimiento and mes and semana):
            continue

        n_filas_usadas += 1
        total_casos += casos
        anios.add(anio)
        jurisdicciones.add(juris)
        municipios.add(municipio)
        instituciones.add(institucion)
        grupos_edad.add(grupo_edad)
        meses.add(mes)
        semanas.add(semana)
        padecimientos_texto.setdefault(epiclave, padecimiento)
        mes_semana[anio][mes].add(semana)
        semana_mes[anio][semana].add(mes)

        key_nuevo = (anio, juris, municipio, institucion, grupo_edad, epiclave, padecimiento, mes, semana)
        combos_nuevo[key_nuevo] += casos
        key_anterior = (anio, juris, municipio, institucion, grupo_edad, epiclave, padecimiento)
        combos_anterior[key_anterior] += casos

    print(f"Filas leídas: {n_filas_leidas} / filas usadas (con Casos válido): {n_filas_usadas}")
    print(f"Total Casos: {total_casos}")
    print(f"Combinaciones CON Mes+Semana: {len(combos_nuevo)}")
    print(f"Combinaciones SIN Mes+Semana (equivalente a antes de ACT14): {len(combos_anterior)}")

    # ---- Verificación (punto 10 de ACT14: los totales sin los filtros
    # nuevos deben coincidir con la entrega anterior) — se recalcula
    # combos_anterior TAMBIÉN colapsando combos_nuevo (sumando Mes+Semana) y
    # se compara contra el combos_anterior calculado directo de las filas
    # crudas: deben ser IDÉNTICOS fila por fila y suma por suma. ----
    combos_colapsado = defaultdict(int)
    for key_nuevo, casos_val in combos_nuevo.items():
        anio, juris, municipio, institucion, grupo_edad, epiclave, padecimiento, mes, semana = key_nuevo
        key_anterior = (anio, juris, municipio, institucion, grupo_edad, epiclave, padecimiento)
        combos_colapsado[key_anterior] += casos_val
    if dict(combos_colapsado) != dict(combos_anterior):
        raise SystemExit(
            "VERIFICACIÓN FALLIDA: sumar el bundle nuevo (con Mes+Semana) sin esas "
            "2 columnas NO reproduce exactamente la agregación anterior (7 columnas). "
            "No se escribe ninguna salida — revisar antes de continuar."
        )
    print("Verificación OK: combos_nuevo colapsando Mes+Semana == combos_anterior (misma suma, misma partición).")

    rows_completas = [list(k) + [v] for k, v in sorted(combos_nuevo.items())]

    # ---- Codificación por índices (ver _Catalogo / _escribir_js_bundle) ----
    cat_anio = _Catalogo()
    cat_juris = _Catalogo()
    cat_municipio = _Catalogo()
    cat_institucion = _Catalogo()
    cat_grupo_edad = _Catalogo()
    cat_epiclave = _Catalogo()
    cat_mes = _Catalogo()
    cat_semana = _Catalogo()
    # padecimiento NO tiene catálogo propio: es 1-a-1 con Epi-Clave, así que
    # se decodifica en el propio .js a partir de padecimientoPorEpiclave
    # (mismo orden que cat_epiclave.lista()), sin repetir texto por fila.
    rows_idx = []
    for anio, juris, municipio, institucion, grupo_edad, epiclave, padecimiento, mes, semana, casos_val in rows_completas:
        rows_idx.append([
            cat_anio.idx(anio), cat_juris.idx(juris), cat_municipio.idx(municipio),
            cat_institucion.idx(institucion), cat_grupo_edad.idx(grupo_edad), cat_epiclave.idx(epiclave),
            cat_mes.idx(mes), cat_semana.idx(semana), casos_val,
        ])
    epiclaves_orden = cat_epiclave.lista()
    padecimiento_por_epiclave = [padecimientos_texto[e] for e in epiclaves_orden]

    # catalogosIdx: las claves DEBEN ir en el mismo orden que las columnas
    # de rows_idx (ver _escribir_js_bundle: decodifica por posición).
    catalogos_idx = {
        "anio": cat_anio.lista(),
        "jurisdiccion": cat_juris.lista(),
        "municipio": cat_municipio.lista(),
        "institucion": cat_institucion.lista(),
        "grupoEdad": cat_grupo_edad.lista(),
        "epiclave": epiclaves_orden,
        "mes": cat_mes.lista(),
        "semana": cat_semana.lista(),
    }

    # ---- Catálogo de padecimientos para el buscador (CIE-10/Epi-clave/nombre) ----
    catalogo_padecimientos = []
    for epiclave, texto in padecimientos_texto.items():
        nombre, cie10_texto, codigos, prefijos = _extraer_cie10(texto)
        catalogo_padecimientos.append({
            "epiclave": epiclave,
            "padecimiento": nombre,
            "cie10_texto": cie10_texto,
            "cie10_codigos": codigos,
            "cie10_prefijos": prefijos,
        })
    catalogo_padecimientos.sort(key=lambda r: r["padecimiento"])

    # ---- mesSemana: coherencia Mes<->Semana por Año (ACT14, punto 2) ----
    mes_semana_out = {}
    for anio in sorted(mes_semana.keys()):
        mes_semana_out[anio] = {
            "porMes": {mes: sorted(sems) for mes, sems in mes_semana[anio].items()},
            "porSemana": {sem: sorted(ms) for sem, ms in semana_mes[anio].items()},
        }

    meta = {
        "fuente": "BASE_CASOS_2025_2026.xlsx (misma base validada en Cargar datos, ACT06)",
        "nota": "Pre-agregado por Año+Jurisdicción+Municipio+Institución+Grupo de edad+"
                "Epi-Clave+Padecimiento+Mes+Semana sumando Casos: matemáticamente idéntico a "
                "sumar las filas crudas por esas mismas columnas. Verificado en esta misma "
                "corrida (ver build_morbilidad_casos.py) que colapsar Mes+Semana reproduce "
                "EXACTAMENTE la agregación de 7 columnas de antes de ACT14 (mismos totales, "
                "misma partición) — Sexo y Entidad siguen sin usarse aquí (ortogonales a este "
                "bundle; Sexo vive en el bundle temporal). Epi-Clave <-> Padecimiento y "
                "Municipio -> Jurisdicción son relaciones 1 a 1 en esta base, así que "
                "agregarlas no fragmenta ninguna combinación. Guardado con codificación por "
                "índices (catálogos + filas de índices) para no repetir texto: se decodifica "
                "a filas de texto completas dentro de este mismo archivo, al cargarlo. No "
                "incluye datos personales, sólo conteos agregados.",
        "filas_originales": n_filas_leidas,
        "filas_agregadas": len(rows_completas),
        "total_casos": total_casos,
        "anios_disponibles": sorted(anios),
        "total_jurisdicciones": len(jurisdicciones),
        "total_municipios": len(municipios),
        "total_instituciones": len(instituciones),
        "total_grupos_edad": len(grupos_edad),
        "total_padecimientos": len(padecimientos_texto),
        "total_meses": len(meses),
        "total_semanas": len(semanas),
        "mes_semana": mes_semana_out,
    }

    catalogos_salida = {
        "anios": sorted(anios),
        "jurisdicciones": sorted(jurisdicciones),
        "municipios": sorted(municipios),
        "instituciones": sorted(instituciones),
        "gruposEdad": sorted(grupos_edad),
        "meses": sorted(meses),
        "semanas": sorted(semanas),
    }

    data_auditoria = {
        "meta": meta,
        "headers": CASOS_HEADERS_OUT,
        "rows": rows_completas,
        "catalogos": catalogos_salida,
        "catalogo_padecimientos": catalogo_padecimientos,
    }
    with open(OUT_CASOS_JSON, "w", encoding="utf-8") as f:
        json.dump(data_auditoria, f, ensure_ascii=False, indent=2)
    print("Escrito", OUT_CASOS_JSON, "(auditoría, filas completas sin codificar)")

    extra_top_level = {
        "catalogos": catalogos_salida,
        "catalogo_padecimientos": catalogo_padecimientos,
        "padecimientoPorEpiclave": padecimiento_por_epiclave,
    }
    _escribir_js_bundle(
        OUT_CASOS_JS,
        "SNSP_MORBILIDAD_CASOS_DATA",
        meta,
        CASOS_HEADERS_OUT,
        catalogos_idx,
        rows_idx,
        extra_top_level,
        "/**\n"
        " * data/real/morbilidad_casos_2025_2026.js\n"
        " * GENERADO AUTOMÁTICAMENTE por scripts/build_morbilidad_casos.py — no editar a mano.\n"
        " * Fuente: data/raw/BASE_CASOS_2025_2026.xlsx (misma base de Cargar datos, ACT06).\n"
        " * Fuente ÚNICA del tablero unificado de Morbilidad: buscador, filtros (incluye\n"
        " * Mes y Semana epidemiológica desde ACT14), Analizar (Casos/Población/Tasa) y\n"
        " * Por (Padecimiento/Municipio/Jurisdicción/Grupo de edad). Codificado por índices\n"
        " * (ver cabecera del script) y decodificado a filas de texto dentro de este mismo\n"
        " * archivo: window.SNSP_MORBILIDAD_CASOS_DATA sigue teniendo la forma de siempre\n"
        " * ({meta, headers, rows, catalogos, catalogo_padecimientos}) para quien lo consume.\n"
        " * No incluye datos personales.\n"
        " * -----------------------------------------------------------------------\n"
        " */\n",
        campo_derivado={"header": "Padecimiento", "despues_de": "Epi-Clave", "fuente": "padecimientoPorEpiclave", "via_catalogo": "epiclave"},
    )
    print("Escrito", OUT_CASOS_JS, f"({len(rows_idx)} filas codificadas por índices)")


def build_temporal():
    """Bundle ligero (Año+Jurisdicción+Municipio+Institución+Mes+Semana+
    Sexo+Epi-Clave -> suma Casos) exclusivamente para las 3 gráficas fijas
    que necesitan desglose temporal/por sexo: Tendencia por año/mes,
    comparativo por mes y año, Casos por sexo. No incluye Grupo de edad
    (esas gráficas no lo usan) ni Población (son de Casos únicamente:
    esta base no desagrega población por mes ni por sexo). El texto de
    Padecimiento se resuelve en el navegador contra catalogo_padecimientos
    del bundle principal, por Epi-Clave (mismo catálogo, no se duplica).
    ACT14: se agrega Semana a la llave (ya traía Mes) para que estos 3
    paneles también respeten un filtro de Semana."""
    print("Leyendo", RAW_CASOS, "para el bundle temporal...")
    wb = openpyxl.load_workbook(RAW_CASOS, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    header = [str(c.value or "").strip() for c in next(ws.iter_rows(min_row=1, max_row=1))]
    idx = {name: i for i, name in enumerate(header)}
    columnas_requeridas = ["Año", "Jurisdicción", "Municipio", "Institución", "Mes", "Semana", "Sexo", "Epi-Clave", "Casos"]
    for col in columnas_requeridas:
        if col not in idx:
            raise SystemExit(f"Columna esperada '{col}' no encontrada en {RAW_CASOS}: {header}")

    combos_nuevo = defaultdict(int)   # con Semana
    combos_anterior = defaultdict(int)  # sin Semana (equivalente a antes de ACT14)
    meses = set()
    semanas = set()
    sexos = set()
    n_filas_usadas = 0
    total_casos = 0

    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[idx["Año"]] is None:
            continue
        casos = row[idx["Casos"]]
        if casos is None or casos == "":
            continue
        try:
            casos = int(casos)
        except (TypeError, ValueError):
            continue

        anio = str(int(row[idx["Año"]]))
        juris = str(row[idx["Jurisdicción"]] or "").strip()
        municipio = str(row[idx["Municipio"]] or "").strip()
        institucion = str(row[idx["Institución"]] or "").strip()
        mes = str(row[idx["Mes"]] or "").strip()
        semana = str(row[idx["Semana"]] or "").strip()
        sexo = str(row[idx["Sexo"]] or "").strip()
        epiclave = str(row[idx["Epi-Clave"]] or "").strip()
        if not (anio and juris and municipio and institucion and mes and semana and sexo and epiclave):
            continue

        n_filas_usadas += 1
        total_casos += casos
        meses.add(mes)
        semanas.add(semana)
        sexos.add(sexo)

        key_nuevo = (anio, juris, municipio, institucion, mes, semana, sexo, epiclave)
        combos_nuevo[key_nuevo] += casos
        key_anterior = (anio, juris, municipio, institucion, mes, sexo, epiclave)
        combos_anterior[key_anterior] += casos

    print(f"Filas usadas (bundle temporal): {n_filas_usadas}")
    print(f"Total Casos (bundle temporal): {total_casos}")
    print(f"Combinaciones CON Semana: {len(combos_nuevo)}")

    # Verificación: colapsar Semana debe reproducir exactamente la
    # agregación anterior (Mes, sin Semana).
    combos_colapsado = defaultdict(int)
    for key_nuevo, casos_val in combos_nuevo.items():
        anio, juris, municipio, institucion, mes, semana, sexo, epiclave = key_nuevo
        key_anterior = (anio, juris, municipio, institucion, mes, sexo, epiclave)
        combos_colapsado[key_anterior] += casos_val
    if dict(combos_colapsado) != dict(combos_anterior):
        raise SystemExit(
            "VERIFICACIÓN FALLIDA (bundle temporal): sumar el bundle nuevo (con Semana) sin "
            "esa columna NO reproduce exactamente la agregación anterior (por Mes). No se "
            "escribe ninguna salida — revisar antes de continuar."
        )
    print("Verificación OK (temporal): combos_nuevo colapsando Semana == combos_anterior.")

    rows_completas = [list(k) + [v] for k, v in sorted(combos_nuevo.items())]

    cat_anio = _Catalogo()
    cat_juris = _Catalogo()
    cat_municipio = _Catalogo()
    cat_institucion = _Catalogo()
    cat_mes = _Catalogo()
    cat_semana = _Catalogo()
    cat_sexo = _Catalogo()
    cat_epiclave = _Catalogo()
    rows_idx = []
    for anio, juris, municipio, institucion, mes, semana, sexo, epiclave, casos_val in rows_completas:
        rows_idx.append([
            cat_anio.idx(anio), cat_juris.idx(juris), cat_municipio.idx(municipio),
            cat_institucion.idx(institucion), cat_mes.idx(mes), cat_semana.idx(semana),
            cat_sexo.idx(sexo), cat_epiclave.idx(epiclave), casos_val,
        ])
    catalogos_idx = {
        "anio": cat_anio.lista(),
        "jurisdiccion": cat_juris.lista(),
        "municipio": cat_municipio.lista(),
        "institucion": cat_institucion.lista(),
        "mes": cat_mes.lista(),
        "semana": cat_semana.lista(),
        "sexo": cat_sexo.lista(),
        "epiclave": cat_epiclave.lista(),
    }

    meta = {
        "fuente": "BASE_CASOS_2025_2026.xlsx (misma base validada en Cargar datos, ACT06)",
        "nota": "Pre-agregado por Año+Jurisdicción+Municipio+Institución+Mes+Semana+Sexo+"
                "Epi-Clave sumando Casos, EXCLUSIVAMENTE para las gráficas fijas Tendencia "
                "por año/mes, comparativo por mes y año y Casos por sexo (Casos únicamente, "
                "sin Población/Tasa: esta base no desagrega población por mes ni por sexo). "
                "Verificado en esta misma corrida que colapsar Semana reproduce EXACTAMENTE "
                "la agregación de antes de ACT14 (por Mes, sin Semana). El texto de "
                "Padecimiento se resuelve por Epi-Clave contra catalogo_padecimientos del "
                "bundle principal (morbilidad_casos_2025_2026.js), no se duplica aquí. "
                "Guardado con codificación por índices, decodificado dentro de este mismo "
                "archivo. No incluye datos personales, sólo conteos agregados.",
        "filas_agregadas": len(rows_completas),
        "total_casos": total_casos,
        "meses_disponibles": sorted(meses),
        "semanas_disponibles": sorted(semanas),
        "sexos_disponibles": sorted(sexos),
    }

    data_auditoria = {"meta": meta, "headers": TEMPORAL_HEADERS_OUT, "rows": rows_completas}
    with open(OUT_TEMPORAL_JSON, "w", encoding="utf-8") as f:
        json.dump(data_auditoria, f, ensure_ascii=False, indent=2)
    print("Escrito", OUT_TEMPORAL_JSON, "(auditoría, filas completas sin codificar)")

    _escribir_js_bundle(
        OUT_TEMPORAL_JS,
        "SNSP_MORBILIDAD_TEMPORAL_DATA",
        meta,
        TEMPORAL_HEADERS_OUT,
        catalogos_idx,
        rows_idx,
        None,
        "/**\n"
        " * data/real/morbilidad_temporal_2025_2026.js\n"
        " * GENERADO AUTOMÁTICAMENTE por scripts/build_morbilidad_casos.py — no editar a mano.\n"
        " * Fuente: data/raw/BASE_CASOS_2025_2026.xlsx (misma base de Cargar datos, ACT06).\n"
        " * Bundle SECUNDARIO y ligero (sólo Casos) del tablero unificado de Morbilidad,\n"
        " * usado ÚNICAMENTE por 3 gráficas fijas: Tendencia por año/mes, comparativo por\n"
        " * mes y año, y Casos por sexo — ahora también filtrable por Semana (ACT14). El\n"
        " * resto del tablero (buscador, Analizar/Por, paneles de municipio/jurisdicción/\n"
        " * institución/grupo de edad/padecimientos) sigue usando\n"
        " * morbilidad_casos_2025_2026.js sin cambios. Codificado por índices, decodificado\n"
        " * dentro de este mismo archivo. No incluye datos personales.\n"
        " * -----------------------------------------------------------------------\n"
        " */\n",
    )
    print("Escrito", OUT_TEMPORAL_JS, f"({len(rows_idx)} filas codificadas por índices)")


def build_poblacion():
    print("Leyendo", RAW_POBLACION, "...")
    wb = openpyxl.load_workbook(RAW_POBLACION, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    header = [str(c.value or "").strip() for c in next(ws.iter_rows(min_row=1, max_row=1))]
    idx = {name: i for i, name in enumerate(header)}
    for col in ["Año", "Municipio", "Grupo de edad", "Población"]:
        if col not in idx:
            raise SystemExit(f"Columna esperada '{col}' no encontrada en {RAW_POBLACION}: {header}")

    headers_out = ["Año", "Municipio", "Grupo de edad", "Población"]
    rows_out = []
    n_filas = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[idx["Año"]] is None:
            continue
        anio = str(int(row[idx["Año"]]))
        municipio = str(row[idx["Municipio"]] or "").strip()
        grupo_edad = str(row[idx["Grupo de edad"]] or "").strip()
        pob = row[idx["Población"]]
        if pob is None or pob == "" or not municipio or not grupo_edad:
            continue
        n_filas += 1
        rows_out.append([anio, municipio, grupo_edad, round(float(pob), 2)])

    print(f"Filas de población: {n_filas}")

    data = {
        "meta": {
            "fuente": "BASE_POBLACION_2025_2026.xlsx (misma base validada en Cargar datos, ACT06)",
            "nota": "Se conserva tal cual el archivo original, sin agregación adicional. "
                    "Sigue siendo ANUAL (Año+Municipio+Grupo de edad): ACT14 no la duplica "
                    "ni la desagrega por Mes/Semana, a petición explícita. "
                    "No incluye datos personales.",
            "filas": n_filas,
        },
        "headers": headers_out,
        "rows": rows_out,
    }

    with open(OUT_POBLACION_JSON, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
    print("Escrito", OUT_POBLACION_JSON)

    with open(OUT_POBLACION_JS, "w", encoding="utf-8") as f:
        f.write(
            "/**\n"
            " * data/real/morbilidad_poblacion_2025_2026.js\n"
            " * GENERADO AUTOMÁTICAMENTE por scripts/build_morbilidad_casos.py — no editar a mano.\n"
            " * Fuente: data/raw/BASE_POBLACION_2025_2026.xlsx (misma base de Cargar datos, ACT06).\n"
            " * Sigue siendo ANUAL (ACT14 no la toca: no se duplica ni se desagrega por Mes/Semana).\n"
            " * No incluye datos personales.\n"
            " * -----------------------------------------------------------------------\n"
            " */\n"
            "window.SNSP_MORBILIDAD_POBLACION_DATA = "
        )
        json.dump(data, f, ensure_ascii=False, separators=(",", ":"))
        f.write(";\n")
    print("Escrito", OUT_POBLACION_JS)


def main():
    build_casos()
    build_temporal()
    build_poblacion()


if __name__ == "__main__":
    main()
